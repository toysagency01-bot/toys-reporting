"""Build a local-only D1 import for AM Klinika leads.

The generated SQL contains client PII and is therefore written only below
`.wrangler/tmp`, which is excluded from version control. The script prints
counts only and never prints lead values.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ID = "prj_amklinika_k8m2"
META_TABS = (
    "Lead_meta_new",
    "Lead_meta_ru",
    "Lead_meta_ru_mechanic",
    "Lead_meta_cz_mechanic",
)
PROJECT_TABS = (
    ("meta", "План работы (Meta)", "План работ", "plan"),
    ("meta", "Еженедельная сводка (Meta)", "Еженедельная сводка", "weekly-report"),
    ("meta", "Месячная сводка (Meta)", "Месячная сводка", "weekly-report"),
    ("meta", "Анализ конкурентов (Meta)", "Анализ конкурентов", "competitors"),
    ("meta", "Креативный бриф (Meta)", "Креативный бриф", "creative-brief"),
    ("meta", "Стратегия (Meta)", "Стратегия", "table"),
    ("project", "План работы (Google)", "План работ", "plan"),
    ("project", "Еженедельная сводка (Google)", "Еженедельная сводка", "weekly-report"),
    ("project", "Месячная сводка (Google)", "Месячная сводка", "weekly-report"),
    ("project", "Ключевые слова (Google)", "Ключевые слова", "table"),
    ("project", "Объявления (Google)", "Объявления", "table"),
)


def sql(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def digest(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def iso_timestamp(value) -> str:
    if isinstance(value, datetime):
        candidate = value
        if candidate.tzinfo is None:
            candidate = candidate.replace(tzinfo=timezone.utc)
        return candidate.isoformat().replace("+00:00", "Z")
    if isinstance(value, date):
        return f"{value.isoformat()}T00:00:00Z"
    candidate = text(value)
    if not candidate or candidate in {"0", "0.0"}:
        return ""
    normalized = candidate.replace(" ", "T")
    try:
        parsed = datetime.fromisoformat(normalized.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.isoformat().replace("+00:00", "Z")
    except ValueError:
        return candidate


def is_test(values: list[str]) -> bool:
    haystack = " ".join(values).casefold()
    markers = (" test ", "тест", "example.com", "reaper", "asdf")
    padded = f" {haystack} "
    return any(marker in padded for marker in markers)


def meta_language(tab: str) -> str:
    if "_ru" in tab:
        return "RU"
    if "_cz" in tab:
        return "CZ"
    return "CZ"


def meta_service(tab: str) -> str:
    return "Автомеханика" if "mechanic" in tab else "Кузовной ремонт"


def row_map(header: list[str], values: tuple) -> dict[str, str]:
    return {key: text(values[index] if index < len(values) else "") for index, key in enumerate(header)}


def lead_statement(lead: dict[str, object]) -> list[str]:
    metadata = json.dumps(lead["metadata"], ensure_ascii=False, separators=(",", ":"))
    values = (
        sql(lead["id"]), sql(PROJECT_ID), sql(lead["source"]), sql(lead["external"]),
        sql(lead["created_at"]), sql(lead.get("name")), sql(lead.get("phone")),
        sql(lead.get("email")), sql(metadata),
    )
    insert = (
        "INSERT INTO leads (id, project_id, source_provider, external_lead_id, created_at, "
        "name, phone, email, metadata_json) VALUES ("
        + ",".join(values)
        + ") ON CONFLICT(project_id, source_provider, external_lead_id) DO UPDATE SET "
        "created_at=excluded.created_at,name=excluded.name,phone=excluded.phone,email=excluded.email,"
        "metadata_json=excluded.metadata_json,imported_at=datetime('now'),updated_at=datetime('now');"
    )
    event_id = f"evt_{digest(str(lead['id']) + ':initial')[:24]}"
    event = (
        "INSERT INTO lead_status_events (id, project_id, lead_id, from_status, to_status, changed_at, note) VALUES ("
        f"{sql(event_id)},{sql(PROJECT_ID)},{sql(lead['id'])},NULL,'Новый',{sql(lead['created_at'])},"
        "'Imported from project spreadsheet') ON CONFLICT(id) DO NOTHING;"
    )
    return [insert, event]


def site_leads(workbook) -> list[dict[str, object]]:
    sheet = workbook["Lead Site"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [text(value) for value in rows[0]]
    expected = ["Дата", "Имя", "Фамилия", "Телефон", "Услуга", "Описание", "Язык", "Страница"]
    if header[: len(expected)] != expected:
        raise ValueError("Lead Site: unexpected header")
    output = []
    for raw in rows[1:]:
        data = row_map(header, raw)
        if not any(data.values()):
            continue
        full_name = " ".join(filter(None, (data.get("Имя"), data.get("Фамилия")))).strip()
        seed = "\x1f".join(text(data.get(key)) for key in expected)
        external = digest(f"Lead Site\x1f{seed}")[:32]
        phone = data.get("Телефон", "")
        check = [full_name, phone, data.get("Описание", "")]
        phone_digits = "".join(character for character in phone if character.isdigit())
        invalid_czech_phone = phone.strip().startswith("+420") and len(phone_digits) != 12
        placeholder_name = full_name.casefold() in {"john doe", "jane doe"}
        output.append({
            "id": f"lead_amk_{digest('site:' + external)[:24]}",
            "source": "google_sheets_site",
            "external": external,
            "created_at": iso_timestamp(data.get("Дата")),
            "name": full_name,
            "phone": phone,
            "email": "",
            "metadata": {
                "sourceTab": "Lead Site",
                "service": data.get("Услуга", ""),
                "description": data.get("Описание", ""),
                "language": data.get("Язык", ""),
                "page": data.get("Страница", ""),
                "isTest": is_test(check) or invalid_czech_phone or placeholder_name,
            },
        })
    return output


def meta_leads(workbook) -> list[dict[str, object]]:
    output = []
    for tab in META_TABS:
        sheet = workbook[tab]
        rows = list(sheet.iter_rows(values_only=True))
        header_index = next(
            (index for index, row in enumerate(rows) if text(row[0] if row else "") == "id"),
            None,
        )
        if header_index is None:
            raise ValueError(f"{tab}: expected header not found")
        header = [text(value) for value in rows[header_index]][:17]
        if header[:2] != ["id", "created_time"]:
            raise ValueError(f"{tab}: invalid header")
        # Exported Meta tabs can place fresh rows above the header. Read every
        # non-header row and keep only the canonical first 17 columns.
        candidates = rows[:header_index] + rows[header_index + 1 :]
        for raw in candidates:
            data = row_map(header, raw[:17])
            external = data.get("id", "")
            if not external or external == "id" or not data.get("created_time"):
                continue
            name = data.get("full_name", "")
            phone = data.get("phone_number", "").removeprefix("p:")
            email = data.get("email", "")
            question_key = next((key for key in header if key not in {
                "id", "created_time", "ad_id", "ad_name", "adset_id", "adset_name",
                "campaign_id", "campaign_name", "form_id", "form_name", "is_organic",
                "platform", "full_name", "phone_number", "email", "lead_status",
            }), "")
            description = data.get(question_key, "") if question_key else ""
            output.append({
                "id": f"lead_amk_{digest('meta:' + external)[:24]}",
                "source": "meta_ads",
                "external": external,
                "created_at": iso_timestamp(data.get("created_time")),
                "name": name,
                "phone": phone,
                "email": email,
                "metadata": {
                    "sourceTab": tab,
                    "service": meta_service(tab),
                    "description": description,
                    "language": meta_language(tab),
                    "campaignId": data.get("campaign_id", ""),
                    "campaign": data.get("campaign_name", ""),
                    "adsetId": data.get("adset_id", ""),
                    "adset": data.get("adset_name", ""),
                    "adId": data.get("ad_id", ""),
                    "ad": data.get("ad_name", ""),
                    "formId": data.get("form_id", ""),
                    "form": data.get("form_name", ""),
                    "platform": data.get("platform", ""),
                    "sourceStatus": data.get("lead_status", ""),
                    "isOrganic": data.get("is_organic", "").casefold() == "true",
                    "isTest": is_test([name, phone, email, description]),
                },
            })
    return output


def worksheet_gviz(sheet) -> str:
    rows = list(sheet.iter_rows(values_only=True))
    width = max((len(row) for row in rows), default=0)
    cols = [{"id": str(index), "label": "", "type": "string"} for index in range(width)]
    output_rows = []
    for row in rows:
        cells = []
        for value in row:
            if value is None:
                cells.append(None)
            elif isinstance(value, (datetime, date)):
                cells.append({"v": iso_timestamp(value), "f": value.isoformat()})
            elif isinstance(value, bool):
                cells.append({"v": value})
            elif isinstance(value, (int, float)):
                cells.append({"v": value})
            else:
                cells.append({"v": str(value)})
        cells.extend([None] * (width - len(cells)))
        output_rows.append({"c": cells})
    return json.dumps(
        {"version": "0.6", "status": "ok", "table": {"cols": cols, "rows": output_rows}},
        ensure_ascii=False,
        separators=(",", ":"),
    )


def project_tab_statements(workbook) -> list[str]:
    statements = []
    for position, (channel, title, label, mode) in enumerate(PROJECT_TABS):
        if title not in workbook.sheetnames:
            continue
        content = worksheet_gviz(workbook[title])
        source_hash = digest(content)
        tab_id = f"tab_amklinika_{digest(title)[:20]}"
        statements.append(
            "INSERT INTO project_tabs (id, project_id, channel, tab_key, source_title, label, mode, "
            "position, content_json, raw_content_json, source_hash) VALUES ("
            f"{sql(tab_id)},{sql(PROJECT_ID)},{sql(channel)},{sql(digest(title)[:24])},{sql(title)},"
            f"{sql(label)},{sql(mode)},{position},'','',{sql(source_hash)}) "
            "ON CONFLICT(project_id, tab_key) DO UPDATE SET label=excluded.label,mode=excluded.mode,"
            "position=excluded.position,content_json='',raw_content_json='',"
            "source_hash=excluded.source_hash,imported_at=datetime('now'),updated_at=datetime('now');"
        )
        # Keep every D1 statement comfortably below SQLite's remote importer
        # limit. The final JSON is byte-for-byte identical to `content`.
        for start in range(0, len(content), 30000):
            chunk = content[start : start + 30000]
            statements.append(
                f"UPDATE project_tabs SET content_json = content_json || {sql(chunk)} WHERE id = {sql(tab_id)};"
            )
        statements.append(
            f"UPDATE project_tabs SET raw_content_json=content_json,source_hash={sql(source_hash)},"
            f"imported_at=datetime('now'),updated_at=datetime('now') WHERE id={sql(tab_id)};"
        )
    return statements


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / ".wrangler" / "tmp" / "amklinika-leads-import.sql",
    )
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    leads = site_leads(workbook) + meta_leads(workbook)
    unique: dict[tuple[str, str], dict[str, object]] = {}
    for lead in leads:
        unique[(str(lead["source"]), str(lead["external"]))] = lead
    # Wrangler's remote D1 file importer wraps the upload atomically and rejects
    # explicit BEGIN/COMMIT statements.
    statements = ["PRAGMA foreign_keys = ON;"]
    tab_statements = project_tab_statements(workbook)
    imported_tab_count = sum(1 for _, title, _, _ in PROJECT_TABS if title in workbook.sheetnames)
    statements.extend(tab_statements)
    for lead in unique.values():
        statements.extend(lead_statement(lead))
    statements.append(
        "INSERT INTO sync_runs (project_id, provider, job_type, status, started_at, finished_at, "
        "rows_read, rows_written, message, details_json) VALUES ("
        f"{sql(PROJECT_ID)},'google_sheets','amklinika-historical-leads','success',datetime('now'),datetime('now'),"
        f"{len(leads)},{len(unique)},'AM Klinika historical leads imported into D1',"
        f"{sql(json.dumps({'leads': len(unique), 'testRowsHidden': sum(bool(item['metadata']['isTest']) for item in unique.values()), 'projectTabs': imported_tab_count}, separators=(',', ':')))}" 
        ");"
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text("\n".join(statements) + "\n", encoding="utf-8")
    print(json.dumps({
        "output": str(args.output),
        "rowsRead": len(leads),
        "leads": len(unique),
        "testRowsHidden": sum(bool(item["metadata"]["isTest"]) for item in unique.values()),
        "projectTabs": imported_tab_count,
    }))


if __name__ == "__main__":
    main()
